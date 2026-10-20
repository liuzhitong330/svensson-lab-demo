"""Reproduce the bounded Piening 2024 reference-data comparison.

Requires Python 3.10+ and openpyxl. Run:
  python analysis/cancer_analysis.py --source /path/to/source.xlsx
or omit --source to download the public source into an OS temporary directory.

The input workbook is never edited. Source rows are not joined across panels.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import statistics
import tempfile
import urllib.request
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-47359-5/MediaObjects/41467_2024_47359_MOESM4_ESM.xlsx"
SHA256 = "2c83fcc4a3232632b88534c22a129cb525210a6314deba29d63f93d90dcf1f4d"
GROUPS = [
    ("nc", "NC", False),
    ("nc_icb", "NC + ICB", True),
    ("wd", "WD", False),
    ("wd_icb", "WD + ICB", True),
    ("sema", "WD + Sema", False),
    ("sema_icb", "WD + Sema + ICB", True),
]
EXPECTED = {
    "tumor_day16": [18, 7, 19, 14, 35, 35],
    "immune": [9, 3, 9, 7, 17, 19],
}


def quantile(values, q):
    ordered = sorted(values)
    p = (len(ordered) - 1) * q
    lo, hi = math.floor(p), math.ceil(p)
    return ordered[lo] + (ordered[hi] - ordered[lo]) * (p - lo)


def summary(values):
    return {
        "n": len(values),
        "mean": statistics.mean(values),
        "sd": statistics.stdev(values) if len(values) > 1 else None,
        "median": statistics.median(values),
        "q1": quantile(values, .25),
        "q3": quantile(values, .75),
    }


def is_numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def extract(source):
    actual_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    if actual_sha != SHA256:
        raise ValueError(f"Source checksum changed: {actual_sha}. Review before use.")
    wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    endpoints = {
        "mass_change_week22": {
            "title": "Body-mass change from week 16 to week 22",
            "unit": "%", "panel": "Fig. 3B",
            "biologicalUnit": "Within-row mouse body-mass record",
            "comparisonAllowed": "mean difference in percentage points; no ratio",
            "groups": [],
        },
        "tumor_day16": {
            "title": "Final tumor volume on day 16",
            "unit": "mm³", "panel": "Fig. 3G",
            "biologicalUnit": "One mouse per observation, as reported in the Fig. 3 caption; source labels are panel-local",
            "comparisonAllowed": "difference and ratio of group means, descriptive only",
            "groups": [],
        },
        "immune": {
            "title": "GzmB+ Perforin+ cells among TRP-2-specific CD8+ TIL",
            "unit": "%", "panel": "Fig. 3I",
            "biologicalUnit": "Mouse-level flow-cytometry summary",
            "comparisonAllowed": "difference in percentage points and ratio of group means, descriptive only",
            "groups": [],
        },
    }
    traces = []
    mass_sheet = wb["Fig. 3B"]
    for key, label, first, last, startcol, expected in [
        ("nc", "NC", 6, 15, 3, 10),
        ("wd", "WD", 19, 28, 3, 10),
        ("sema", "WD + Sema", 6, 25, 12, 20),
    ]:
        values, source_cells = [], []
        for row in range(first, last+1):
            grams = [mass_sheet.cell(row, startcol+i).value for i in range(7)]
            assert all(is_numeric(v) and v > 0 for v in grams)
            percent = [100*(v / grams[0]-1) for v in grams]
            cell_range = f"Fig. 3B!{openpyxl.utils.get_column_letter(startcol)}{row}:{openpyxl.utils.get_column_letter(startcol+6)}{row}"
            values.append(percent[-1])
            source_cells.append(cell_range)
            traces.append({
                "key": key, "label": label,
                "sourceLabel": mass_sheet.cell(row, startcol-1).value,
                "sourceRow": row, "sourceRange": cell_range,
                "weeks": [16, 17, 18, 19, 20, 21, 22],
                "grams": grams, "percentChange": percent,
            })
        assert len(values) == expected
        endpoints["mass_change_week22"]["groups"].append({
            "key": key, "label": label, "icb": False,
            "values": values, "sourceCells": source_cells, **summary(values),
        })
    for endpoint, panel, last in [
        ("tumor_day16", "Fig. 3G", 39), ("immune", "Fig. 3I", 23)
    ]:
        sheet = wb[panel]
        for group_index, (key, label, icb) in enumerate(GROUPS):
            col = group_index+3
            assert sheet.cell(4, col).value == label
            values, source_cells, source_labels = [], [], []
            for row in range(5, last+1):
                value = sheet.cell(row, col).value
                if value is None:
                    continue
                assert is_numeric(value) and value >= 0
                values.append(value)
                source_cells.append(f"{panel}!{openpyxl.utils.get_column_letter(col)}{row}")
                source_labels.append(sheet.cell(row, 2).value)
            assert len(values) == EXPECTED[endpoint][group_index]
            endpoints[endpoint]["groups"].append({
                "key": key, "label": label, "icb": icb, "values": values,
                "sourceCells": source_cells, "sourceLabels": source_labels,
                **summary(values),
            })
    result = {
        "study": {
            "doi": "10.1038/s41467-024-47359-5",
            "title": "Obesity-related T cell dysfunction impairs immunosurveillance and increases cancer risk",
            "authors": "Piening et al.",
            "year": 2024,
            "url": "https://www.nature.com/articles/s41467-024-47359-5",
            "ownership": "External reference study, not Svensson Lab data",
            "license": "CC BY 4.0; quantitative transformation and presentation are new",
        },
        "source": {"url": URL, "sha256": SHA256, "bytes": source.stat().st_size},
        "endpoints": endpoints,
        "massTraces": traces,
        "design": {
            "model": "B16 melanoma in C57BL/6 mice after normal chow or Western diet",
            "massPanel": "One representative body-mass experiment; six weeks of semaglutide following 12 weeks of Western diet",
            "tumorImmunePanels": "Two pooled experiments; ICB was anti-PD-1 plus anti-CTLA-4",
            "reportedExperimentalUnit": "The Fig. 3 caption reports each n as an individual mouse. The workbook does not provide cross-panel matching or experiment-batch identifiers.",
            "groupLabels": {"NC": "Normal chow", "WD": "Western diet", "Sema": "Semaglutide", "ICB": "Immune checkpoint blockade"},
        },
        "methods": {
            "massChange": "For each mouse row, 100*(week22/week16 - 1), then summarize the resulting percentages.",
            "summaries": "Arithmetic mean, sample SD (ddof=1), median, and linearly interpolated 25th and 75th percentiles.",
            "missingness": "Preserve recorded numeric zero. Omit blank endpoint cells without replacing them. Display panel-specific sample sizes.",
            "comparison": "Differences/ratios are unadjusted descriptions of groups, not causal estimates, inferential tests, or clinical predictions.",
        },
        "limits": [
            "The body-mass experiment and pooled tumor/immune experiments are distinct panels. Mouse/Tumor numbers are local source labels, not cross-panel identifiers.",
            "Never correlate or pair mass, tumor, and immune measurements across panels.",
            "Fig. 3 reports each n as an individual mouse. Panel-local source labels do not establish cross-panel matching or identify the two pooled experiment batches. No confidence intervals or P values are computed.",
            "ICB sample sizes differ substantially between groups, including only three NC + ICB mouse measurements for the immune endpoint.",
            "Body-mass change has no ICB subdivision in the source. The ICB control must be disabled for this endpoint.",
            "The data show endpoint discordance in this model, not a mechanism. They cannot distinguish direct tumor-cell GLP-1R action from host-metabolic or immune effects.",
            "No comparison with diet-switch cohorts, other tumor models, or human treatment outcomes is computed.",
        ],
        "checks": [
            "All included source values numeric and nonnegative; all 40 mass trajectories have seven positive observations.",
            "Fig. 3B: 10 NC, 10 WD, 20 Sema source rows.",
            "Fig. 3G: 18 NC, 7 NC+ICB, 19 WD, 14 WD+ICB, 35 Sema, 35 Sema+ICB mouse observations, per the figure caption.",
            "Fig. 3I: 9 NC, 3 NC+ICB, 9 WD, 7 WD+ICB, 17 Sema, 19 Sema+ICB mouse summaries.",
        ],
    }
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT/"cancer-data.json")
    args = parser.parse_args()
    if args.source:
        result = extract(args.source)
    else:
        with tempfile.TemporaryDirectory(prefix="svensson-cancer-source-") as tmp:
            path = Path(tmp)/"piening2024.xlsx"
            urllib.request.urlretrieve(URL, path)
            result = extract(path)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False)+"\n")
    print(json.dumps({
        "source_sha256": SHA256,
        "summaries": {
            name: {g["key"]: {"n": g["n"], "mean": g["mean"], "sd": g["sd"], "median": g["median"]} for g in endpoint["groups"]}
            for name, endpoint in result["endpoints"].items()
        },
    }, indent=2))


if __name__ == "__main__":
    main()
